"""
Extracts text from a Word (.docx) file — same role as
extraer_texto_pdf.py, for a base CV uploaded in Word format instead of PDF.
"""
import io

from docx import Document


class ExtraccionDOCXError(Exception):
    """El archivo no se pudo abrir, o no tiene texto extraíble."""


def extraer_texto_docx(contenido_docx: bytes) -> str:
    """
    Recibe los bytes crudos de un .docx y devuelve su texto: los párrafos
    normales más el contenido de cualquier tabla (algunas plantillas de CV
    maquetan con tablas, y perder ese texto sería perder datos reales del
    candidato).
    """
    try:
        documento = Document(io.BytesIO(contenido_docx))
        partes = [parrafo.text for parrafo in documento.paragraphs]
        for tabla in documento.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(celda.text for celda in fila.cells))
    except Exception as exc:  # python-docx can raise several exception types
        raise ExtraccionDOCXError(f"No se pudo leer el archivo Word: {exc}") from exc

    texto = "\n".join(parte for parte in partes if parte.strip()).strip()
    if not texto:
        raise ExtraccionDOCXError("El archivo Word no contiene texto.")
    return texto
